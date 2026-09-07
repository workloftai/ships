#!/usr/bin/env python3
"""office_pack.py — turn structured data into EDITABLE Office files.

An agent that only emits PDFs hands the reader a photograph of a document:
final, read-only, nothing to change. Half the time the person on the other
end wants the opposite. They want to tweak a number, add a row, drop a
paragraph, and send it on. That needs a real .docx or .xlsx, not a render.

This module is the generate-from-data half of that. Feed it a plain dict of
sections and a list of tracker rows; get back a Word document and an Excel
workbook that open natively in Office, Google Docs/Sheets and LibreOffice,
with live formulas and editable text. No Office install, no headless browser,
no COM automation. Just python-docx and openpyxl, which are the right tool
when the job is "build the file", not "edit a file someone else made".

Pair it with OfficeCLI (validate / inspect / render / agent-edit) for the
read side. build_demo.py shows both halves working together.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- brand tokens (neutral, product-agnostic; override per deliverable) ------
NAVY = RGBColor(0x1A, 0x1A, 0x2E)
BLUE = RGBColor(0x25, 0x63, 0xEB)
MUTED = RGBColor(0x6B, 0x72, 0x80)
NAVY_HEX = "1A1A2E"
BLUE_HEX = "2563EB"
OFFWHITE_HEX = "F8F9FA"


# --- document spec -----------------------------------------------------------
@dataclass
class Table:
    """A table inside a document section. First row of `rows` is the header."""
    rows: list[list[str]]


@dataclass
class Section:
    heading: str
    body: list[str] = field(default_factory=list)   # paragraphs
    table: Table | None = None


@dataclass
class DocSpec:
    title: str
    subtitle: str
    prepared_for: str
    footer: str
    sections: list[Section] = field(default_factory=list)


# --- .docx generation --------------------------------------------------------
def _shade(cell, hex_fill: str) -> None:
    """Set a table cell background (python-docx has no direct API for it)."""
    tcPr = cell._tc.get_or_add_tcPr()
    # w:shd requires val (the shading pattern) and color; omitting val fails
    # OpenXML schema validation even though Word renders it fine.
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:val"): "clear",
        qn("w:color"): "auto",
        qn("w:fill"): hex_fill,
    })
    tcPr.append(shd)


def build_docx(spec: DocSpec, path: str) -> str:
    """Render a branded, fully editable Word document from `spec`."""
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Cover-ish header block
    band = doc.add_paragraph()
    run = band.add_run(spec.title)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    r = sub.add_run(spec.subtitle)
    r.font.size = Pt(12)
    r.font.color.rgb = BLUE
    r.font.bold = True

    meta = doc.add_paragraph()
    r = meta.add_run(spec.prepared_for)
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED

    for sec in spec.sections:
        h = doc.add_heading(level=1)
        hr = h.add_run(sec.heading)
        hr.font.color.rgb = NAVY
        hr.font.size = Pt(15)
        for para in sec.body:
            doc.add_paragraph(para)
        if sec.table:
            _add_table(doc, sec.table)

    foot = doc.add_paragraph()
    fr = foot.add_run(spec.footer)
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    fr.italic = True

    doc.save(path)
    return path


def _add_table(doc: Document, table: Table) -> None:
    header, *body = table.rows
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, label in enumerate(header):
        _shade(hdr[i], NAVY_HEX)
        p = hdr[i].paragraphs[0]
        run = p.add_run(str(label))
        run.font.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for ri, row in enumerate(body):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[ci], OFFWHITE_HEX)
            p = cells[ci].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)


# --- .xlsx generation --------------------------------------------------------
def build_tracker(
    headers: list[str],
    rows: list[list],
    path: str,
    *,
    sheet_title: str = "Placements",
    formula_cols: dict[int, str] | None = None,
    total_cols: list[int] | None = None,
) -> str:
    """Render an editable Excel tracker with live formulas and a totals row.

    headers      column labels (row 1)
    rows         data rows; a cell may be a value, or None where a formula fills it
    formula_cols {0-based col index: "=<expr with {r}>"} — {r} is the 1-based
                 spreadsheet row, so the formula stays live and editable in Excel
    total_cols   0-based columns to SUM in a bold totals row at the bottom
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    thin = Side(style="thin", color="E5E7EB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor=NAVY_HEX)
    stripe = PatternFill("solid", fgColor=OFFWHITE_HEX)
    header_font = Font(bold=True, color="FFFFFF", size=10)

    for c, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    formula_cols = formula_cols or {}
    first_data, last_data = 2, len(rows) + 1
    for i, row in enumerate(rows):
        r = first_data + i
        for c, val in enumerate(row):
            if c in formula_cols:
                val = formula_cols[c].format(r=r)
            cell = ws.cell(row=r, column=c + 1, value=val)
            cell.border = border
            if i % 2 == 1:
                cell.fill = stripe

    if total_cols:
        r = last_data + 1
        label_cell = ws.cell(row=r, column=1, value="TOTAL")
        label_cell.font = Font(bold=True, color=NAVY_HEX)
        for c in total_cols:
            col = get_column_letter(c + 1)
            cell = ws.cell(
                row=r, column=c + 1,
                value=f"=SUM({col}{first_data}:{col}{last_data})",
            )
            cell.font = Font(bold=True, color=NAVY_HEX)
            cell.border = border

    # Freeze the header and widen columns to their content
    ws.freeze_panes = "A2"
    for c, label in enumerate(headers, start=1):
        width = max(len(str(label)), 12)
        for row in rows:
            if c - 1 < len(row) and row[c - 1] is not None:
                width = max(width, len(str(row[c - 1])) + 2)
        ws.column_dimensions[get_column_letter(c)].width = min(width, 40)

    wb.save(path)
    return path
